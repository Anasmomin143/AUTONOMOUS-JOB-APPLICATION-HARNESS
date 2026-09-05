"""Excel projection of applications.json (spec §21, §22).

JSON is source of truth; this file is regenerated top-to-bottom on every
update. Users should treat the Excel as read-only while the harness runs.
"""
from __future__ import annotations
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..paths import EXCEL_TRACKER

COLUMNS = [
    "Application ID", "Batch ID", "Company", "Role", "Job URL", "Job Source",
    "Location", "Work Mode", "Match Score", "Job Posted Date", "Date Discovered",
    "Date Prepared", "Date Applied", "Resume Type", "Resume Filename",
    "Application Status", "ATS Status", "Email Status", "Interview Stage",
    "Recruiter", "Recruiter URL", "Follow-up Date", "Last Update",
    "Next Action", "Notes", "Error/Reason",
]

DASHBOARD_LABELS = [
    "Total Jobs", "Qualified Jobs", "Applications Prepared",
    "Applications Submitted", "Applications Failed",
    "Assessments", "Interviews", "Offers", "Rejections",
    "Follow-ups Due",
    "", "Response Rate", "Interview Conversion Rate",
    "Offer Conversion Rate", "Rejection Rate",
]


def _header_style(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _ensure_workbook() -> Workbook:
    if EXCEL_TRACKER.exists():
        return load_workbook(EXCEL_TRACKER)
    wb = Workbook()
    apps = wb.active
    apps.title = "Applications"
    for i, col in enumerate(COLUMNS, start=1):
        c = apps.cell(row=1, column=i, value=col)
        _header_style(c)
        apps.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 2)
    apps.freeze_panes = "A2"
    apps.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    dash = wb.create_sheet("Dashboard")
    for i, label in enumerate(DASHBOARD_LABELS, start=1):
        dash.cell(row=i, column=1, value=label).font = Font(bold=True)
        dash.cell(row=i, column=2, value=0 if label and "Rate" not in label else "")
    dash.column_dimensions["A"].width = 30
    dash.column_dimensions["B"].width = 16
    return wb


def rewrite(applications: Iterable[dict], dashboard: dict) -> None:
    EXCEL_TRACKER.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure_workbook()
    apps = wb["Applications"]

    # Wipe existing rows below header.
    if apps.max_row > 1:
        apps.delete_rows(2, apps.max_row - 1)

    for row_i, a in enumerate(applications, start=2):
        row = [
            a.get("application_id"),
            a.get("batch_id"),
            a.get("company"),
            a.get("role"),
            a.get("job_url"),
            a.get("job_source"),
            a.get("location"),
            a.get("work_mode"),
            a.get("match_score"),
            a.get("job_posted_date"),
            a.get("date_discovered"),
            a.get("date_prepared"),
            a.get("date_applied"),
            a.get("resume_type"),
            a.get("resume_file"),
            a.get("status"),
            a.get("ats_status"),
            a.get("email_status"),
            a.get("interview_stage"),
            a.get("recruiter"),
            a.get("recruiter_url"),
            a.get("followup_date"),
            a.get("last_update"),
            a.get("next_action"),
            a.get("notes"),
            a.get("error_reason"),
        ]
        for col_i, val in enumerate(row, start=1):
            apps.cell(row=row_i, column=col_i, value=val)

    # Dashboard values.
    dash = wb["Dashboard"]
    for i, label in enumerate(DASHBOARD_LABELS, start=1):
        if not label:
            continue
        dash.cell(row=i, column=2, value=dashboard.get(label, 0 if "Rate" not in label else ""))

    wb.save(EXCEL_TRACKER)
